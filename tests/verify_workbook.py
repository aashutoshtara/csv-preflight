"""Independent checks for synthetic workbooks produced by workbook.test.ts.

Usage: python3 tests/verify_workbook.py [tests/output/workbook-fixtures.json]
This reads only local synthetic fixtures; it does not run Excel or claim native
spreadsheet compatibility. Exit nonzero for any fidelity or package failure.
"""

import hashlib
import json
from pathlib import Path
import sys
import xml.etree.ElementTree as ET
import zipfile

try:
    import openpyxl
    from openpyxl.utils.cell import range_boundaries
except ImportError as error:
    raise SystemExit("Independent verification requires the openpyxl Python package.") from error


def check_fixture(base, fixture):
    path = (base / fixture["file"]).resolve()
    if path.parent != base.resolve():
        raise AssertionError("Fixture path must stay inside the output folder")
    expected = fixture["expected"]
    rows, columns = len(expected), len(expected[0])
    with zipfile.ZipFile(path) as archive:
        names = archive.namelist()
        assert len(names) == len(set(names)), "Duplicate package parts"
        assert not any("vbaproject" in name.lower() or "externallinks/" in name.lower()
                       or "activex/" in name.lower() or "embeddings/" in name.lower() for name in names), "Macro or external content"
        for name in names:
            if not name.endswith((".xml", ".rels")):
                continue
            data = archive.read(name)
            assert b"<!DOCTYPE" not in data.upper() and b"<!ENTITY" not in data.upper(), "Unsupported XML declaration"
            xml = ET.fromstring(data)
            for element in xml.iter():
                local = element.tag.split("}")[-1]
                assert local not in ("f", "formula1", "formula2"), "Formula element"
                if local == "Relationship":
                    assert element.attrib.get("TargetMode", "Internal").lower() != "external", "External relationship"
                    assert not element.attrib.get("Type", "").endswith("/hyperlink"), "Hyperlink relationship"
                assert "macroEnabled" not in element.attrib.get("ContentType", ""), "Macro-enabled workbook"

    book = openpyxl.load_workbook(path, data_only=False, keep_links=True)
    assert book.sheetnames == ["CSV Data"], "Unexpected sheet list"
    assert not book._external_links, "External workbook link"
    assert set(book.defined_names) == {"CSV_Data_Range"}, "Changed named ranges"
    destinations = list(book.defined_names["CSV_Data_Range"].destinations)
    assert len(destinations) == 1 and destinations[0][0] == "CSV Data", "Wrong named-range sheet"
    assert range_boundaries(destinations[0][1]) == (1, 1, columns, rows), "Changed declared rectangle"
    sheet = book["CSV Data"]
    assert sheet.max_row == rows and sheet.max_column == columns, "Lost trailing blank dimensions"
    assert not sheet.merged_cells.ranges, "Merged data cells"
    nonempty = 0
    for r, values in enumerate(expected, 1):
        assert len(values) == columns, "Malformed expected fixture"
        for c, wanted in enumerate(values, 1):
            cell = sheet.cell(r, c)
            assert cell.data_type != "f" and cell.hyperlink is None, f"Executable cell at {cell.coordinate}"
            assert cell.number_format == "@", f"Wrong text format at {cell.coordinate}"
            if wanted == "":
                assert cell.value in (None, ""), f"Changed blank at {cell.coordinate}"
            else:
                assert type(cell.value) is str and cell.data_type == "s", f"Non-string at {cell.coordinate}"
                assert cell.value == wanted, f"Changed string at {cell.coordinate}"
                nonempty += 1
    return {
        "file": fixture["file"], "rows": rows, "columns": columns,
        "nonempty_cells": nonempty, "blank_cells": rows * columns - nonempty,
        "sha256": hashlib.sha256(path.read_bytes()).hexdigest(), "passed": True,
    }


def main():
    manifest_path = Path(sys.argv[1]) if len(sys.argv) > 1 else Path(__file__).parent / "output/workbook-fixtures.json"
    manifest = json.loads(manifest_path.read_text())
    results = [check_fixture(manifest_path.parent, fixture) for fixture in manifest["fixtures"]]
    report = {
        "reader": "openpyxl", "reader_version": openpyxl.__version__,
        "scope": "Independent local synthetic fixture verification, not native Excel/Sheets testing",
        "fixtures": results, "passed": len(results), "failed": 0,
    }
    output = manifest_path.parent / "workbook-independent-verification.json"
    output.write_text(json.dumps(report, indent=2) + "\n")
    print(json.dumps(report, indent=2))


if __name__ == "__main__":
    main()
